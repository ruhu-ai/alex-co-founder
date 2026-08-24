"""Document builders (docs/15): spec -> build -> validate.

The model emits DATA (JSON specs); this module owns layout, styles, and the
file format — the model can produce a bad spec, but never a corrupt file.
Errors as data throughout (principle 2). Build to the artifact path only after
validation passes: no partial artifacts.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
import threading
import xml.etree.ElementTree as ET
import zipfile

from services import storage

KINDS = ("docx", "xlsx", "pptx", "pdf")

_MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
}

_REQUIRED_PART = {"docx": "word/document.xml", "xlsx": "xl/workbook.xml",
                  "pptx": "ppt/presentation.xml"}


def mime_for(kind: str) -> str:
    return _MIME[kind]


def spec_hash(spec: dict) -> str:
    return "sha256:" + hashlib.sha256(
        json.dumps(spec, sort_keys=True, default=str).encode()).hexdigest()[:16]


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:40] or "document"


def _err(message: str) -> dict:
    return {"status": "error", "error": True, "message": message}


# ---------------------------------------------------------------------------
# builders
# ---------------------------------------------------------------------------

def build_docx(title: str, spec: dict, path: str) -> dict:
    """spec: {sections: [{heading?, paragraphs?: [...], bullets?: [...]}]}"""
    sections = spec.get("sections")
    if not isinstance(sections, list) or not sections:
        return _err("docx spec needs a non-empty 'sections' list")
    try:
        from docx import Document

        doc = Document()
        doc.add_heading(title, 0)
        for sec in sections:
            if not isinstance(sec, dict):
                return _err("each section must be an object")
            if sec.get("heading"):
                doc.add_heading(str(sec["heading"]), level=1)
            for para in sec.get("paragraphs", []) or []:
                doc.add_paragraph(str(para))
            for bullet in sec.get("bullets", []) or []:
                doc.add_paragraph(str(bullet), style="List Bullet")
        doc.save(path)
        return {"status": "success"}
    except Exception as exc:
        return _err(f"docx build failed: {exc}")


def build_xlsx(spec: dict, path: str) -> dict:
    """spec: {sheets: [{name, columns: [...], rows: [[...]], formulas: {A1: "=…"}}]}"""
    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        return _err("xlsx spec needs a non-empty 'sheets' list")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        for i, sheet in enumerate(sheets):
            if not isinstance(sheet, dict) or not sheet.get("name"):
                return _err("each sheet needs a 'name'")
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = str(sheet["name"])[:31]
            columns = sheet.get("columns") or []
            if columns:
                ws.append([str(c) for c in columns])
                for cell in ws[1]:
                    cell.font = Font(bold=True)
            for row in sheet.get("rows", []) or []:
                ws.append(list(row))
            for cell, formula in (sheet.get("formulas") or {}).items():
                ws[cell] = str(formula)
        wb.save(path)
        return {"status": "success"}
    except Exception as exc:
        return _err(f"xlsx build failed: {exc}")


def build_pptx(title: str, spec: dict, path: str) -> dict:
    """spec: {slides: [{title, bullets: [...], notes?}]}"""
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        return _err("pptx spec needs a non-empty 'slides' list")
    try:
        from pptx import Presentation

        prs = Presentation()
        first = prs.slides.add_slide(prs.slide_layouts[0])
        first.shapes.title.text = title
        if first.placeholders and len(first.placeholders) > 1:
            first.placeholders[1].text = "Produced by Co-Founder"
        for slide in slides:
            if not isinstance(slide, dict) or not slide.get("title"):
                return _err("each slide needs a 'title'")
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = str(slide["title"])
            body = s.placeholders[1].text_frame
            bullets = [str(b) for b in (slide.get("bullets") or [])]
            if bullets:
                body.text = bullets[0]
                for b in bullets[1:]:
                    body.add_paragraph().text = b
            if slide.get("notes"):
                s.notes_slide.notes_text_frame.text = str(slide["notes"])
        prs.save(path)
        return {"status": "success"}
    except Exception as exc:
        return _err(f"pptx build failed: {exc}")


def _soffice() -> str | None:
    import shutil

    return shutil.which("soffice") or shutil.which("libreoffice")


# Server-side conversion posture (docs/15 §LibreOffice): bounded concurrency,
# an input size cap, macro parts stripped from OOXML inputs, a fresh per-job
# user profile (macro security defaults to High = disabled), per-job temp
# storage, a hard timeout, and output validation before anything is served.
_SOFFICE_SEMAPHORE = None  # lazy: created on first use inside any thread
_SEMAPHORE_LOCK = threading.Lock()
MAX_CONVERT_BYTES = 25 * 1024 * 1024  # 25 MB input cap
_CONVERT_TIMEOUT_S = 120
_MACRO_PART = re.compile(r"vbaProject|macroSheet|activeX|oleObject", re.IGNORECASE)


def _semaphore() -> threading.BoundedSemaphore:
    global _SOFFICE_SEMAPHORE
    with _SEMAPHORE_LOCK:
        if _SOFFICE_SEMAPHORE is None:
            _SOFFICE_SEMAPHORE = threading.BoundedSemaphore(2)
    return _SOFFICE_SEMAPHORE


def _strip_ooxml_macros(src: str, dst: str) -> None:
    """Copy an OOXML package without macro-bearing parts. Uploaded Office
    files are semi-untrusted; stripping is the belt to the fresh high-security
    profile's brace (non-OOXML inputs pass through unchanged)."""
    import shutil

    if not zipfile.is_zipfile(src):
        shutil.copyfile(src, dst)
        return
    with zipfile.ZipFile(src) as zin, \
            zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if _MACRO_PART.search(item.filename):
                continue
            zout.writestr(item, zin.read(item.filename))


def convert_to_pdf(src_path: str, out_path: str) -> dict:
    """Any Office doc -> PDF via headless LibreOffice (preview path, docs/15).
    Errors as data where soffice is absent."""
    import shutil
    import subprocess

    soffice = _soffice()
    if not soffice:
        return _err("preview needs LibreOffice (soffice) on this machine")
    try:
        size = os.path.getsize(src_path)
    except OSError as exc:
        return _err(f"source unreadable: {exc}")
    if size > MAX_CONVERT_BYTES:
        return _err(
            f"file too large to preview ({size // (1024 * 1024)} MB > 25 MB)")
    if not _semaphore().acquire(timeout=30):
        return _err("preview service is busy — try again in a moment")
    try:
        # Per-job profile + work dir: no shared state between conversions, no
        # lock-check hacks, no attaching to a developer's GUI instance.
        with tempfile.TemporaryDirectory(prefix="cofounder-soffice-") as job:
            profile_dir = os.path.join(job, "profile")
            work_dir = os.path.join(job, "work")
            os.makedirs(work_dir)
            staged = os.path.join(work_dir, os.path.basename(src_path))
            _strip_ooxml_macros(src_path, staged)
            proc = subprocess.run(
                [soffice, "--headless", "--invisible", "--nologo", "--norestore",
                 "--nodefault",
                 f"-env:UserInstallation=file://{profile_dir}",
                 "--convert-to", "pdf", "--outdir", work_dir, staged],
                capture_output=True, timeout=_CONVERT_TIMEOUT_S)
            converted = os.path.join(
                work_dir, os.path.splitext(os.path.basename(staged))[0] + ".pdf")
            if proc.returncode != 0 or not os.path.exists(converted):
                return _err(f"pdf conversion failed: {proc.stderr.decode()[:200]}")
            os.makedirs(os.path.dirname(out_path), exist_ok=True)
            shutil.move(converted, out_path)  # temp may be another filesystem
        valid = validate_document(out_path, "pdf")
        if valid["status"] != "success":
            try:
                os.unlink(out_path)
            except OSError:
                pass
            return _err(f"converted PDF failed validation: {valid['message']}")
        return {"status": "success"}
    except subprocess.TimeoutExpired:
        return _err(f"pdf conversion timed out ({_CONVERT_TIMEOUT_S}s)")
    except Exception as exc:
        return _err(f"pdf conversion failed: {exc}")
    finally:
        _semaphore().release()


def build_pdf(title: str, spec: dict, path: str) -> dict:
    """PDF = docx build + headless LibreOffice conversion (docs/15). Degrades
    to error-as-data where soffice is absent (e.g. slim deploy images)."""
    if not _soffice():
        return _err("PDF conversion needs LibreOffice (soffice) on this machine — "
                    "produce a .docx instead; it converts in one step")
    docx_tmp = path + ".docx"
    built = build_docx(title, spec, docx_tmp)
    if built["status"] != "success":
        return built
    try:
        result = convert_to_pdf(docx_tmp, path)
        return result
    finally:
        if os.path.exists(docx_tmp):
            os.unlink(docx_tmp)


# ---------------------------------------------------------------------------
# validation gate — nothing ships without passing
# ---------------------------------------------------------------------------

def validate_document(path: str, kind: str) -> dict:
    """OOXML: ZIP integrity + required part + XML well-formedness.
    PDF: header + EOF marker (structural sanity, not full parsing)."""
    if kind == "pdf":
        try:
            with open(path, "rb") as fh:
                head = fh.read(8)
                fh.seek(-1024, os.SEEK_END)
                tail = fh.read()
            if not head.startswith(b"%PDF-"):
                return _err("not a valid PDF (bad header)")
            if b"%%EOF" not in tail:
                return _err("not a valid PDF (missing EOF marker)")
            return {"status": "success"}
        except Exception as exc:
            return _err(f"validation failed: {exc}")
    if not zipfile.is_zipfile(path):
        return _err("not a valid OOXML package (not a zip)")
    try:
        with zipfile.ZipFile(path) as zf:
            bad = zf.testzip()
            if bad:
                return _err(f"corrupt member in package: {bad}")
            names = set(zf.namelist())
            required = _REQUIRED_PART[kind]
            if required not in names:
                return _err(f"missing required part: {required}")
            for name in names:
                if name.endswith((".xml", ".rels")):
                    ET.fromstring(zf.read(name))
    except ET.ParseError as exc:
        return _err(f"malformed XML in package: {exc}")
    except Exception as exc:
        return _err(f"validation failed: {exc}")
    return {"status": "success"}


# ---------------------------------------------------------------------------
# the one entry point: spec -> validated artifact
# ---------------------------------------------------------------------------

def produce(kind: str, title: str, spec: dict, artifact_name: str) -> dict:
    """Build to a temp file, validate, promote to the artifact store."""
    if kind not in KINDS:
        return _err(f"unknown kind '{kind}' — one of {KINDS}")
    if not isinstance(spec, dict):
        return _err("spec must be a JSON object")
    build = {"docx": lambda p: build_docx(title, spec, p),
             "xlsx": lambda p: build_xlsx(spec, p),
             "pptx": lambda p: build_pptx(title, spec, p),
             "pdf": lambda p: build_pdf(title, spec, p)}[kind]
    fd, tmp = tempfile.mkstemp(suffix=f".{kind}")
    os.close(fd)
    try:
        built = build(tmp)
        if built["status"] != "success":
            return built
        valid = validate_document(tmp, kind)
        if valid["status"] != "success":
            return valid
        with open(tmp, "rb") as fh:
            storage.save_bytes(artifact_name, fh.read())
        return {"status": "success", "artifact_name": artifact_name,
                "bytes": os.path.getsize(storage.artifact_path(artifact_name))}
    finally:
        os.unlink(tmp)
