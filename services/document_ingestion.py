"""Safe, structured source-document ingestion and scoped retrieval.

The upload boundary validates bytes before storage. Workers then extract bounded
page/slide/section chunks and persist them under an owner-scoped artifact. Model
content is data only: authorization, status transitions, archive limits, and
profile mutation scope are deterministic code contracts.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import math
import os
import re
import zipfile
from collections.abc import Callable
from pathlib import PurePosixPath, PureWindowsPath
from typing import Any

from services import firestore, storage

EXTRACTOR_VERSION = "alex-document-extractor:v1"
MAX_ARCHIVE_MEMBERS = 5_000
MAX_ARCHIVE_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_MEMBER_BYTES = 25 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MAX_DOCUMENT_UNITS = 250
MAX_CHUNK_CHARS = 6_000
MAX_CHUNKS = 400
MAX_SHEET_ROWS = 10_000
MAX_SHEET_COLUMNS = 200

TERMINAL_STATUSES = {
    "READY", "NEEDS_FOUNDER", "CONFIRMED", "NO_TEXT", "UNSUPPORTED", "FAILED"
}
PROCESSING_STATUSES = {"QUEUED", "VALIDATING", "EXTRACTING", "INDEXING"}

_OOXML = {
    ".docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "word/document.xml",
    ),
    ".pptx": (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "ppt/presentation.xml",
    ),
    ".xlsx": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xl/workbook.xml",
    ),
}
_SIMPLE_MIME = {
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".csv": "text/csv",
}
_SUPPORTED = {**{k: v[0] for k, v in _OOXML.items()}, **_SIMPLE_MIME}
_LEGACY_EXTS = {".doc", ".ppt", ".xls"}
_OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")

EmbedFn = Callable[[list[str]], list[list[float]]]
_embed_fn: EmbedFn | None = None


def set_embed_fn(fn: EmbedFn | None) -> None:
    """Inject the optional embedding backend used for hybrid retrieval."""
    global _embed_fn
    _embed_fn = fn


def _error(code: str, message: str, *, status: str = "FAILED") -> dict[str, Any]:
    return {"status": "error", "error": True, "error_code": code,
            "ingestion_status": status, "message": message}


def _safe_archive_name(name: str) -> bool:
    if not name or name in {".", "./"}:
        return True
    if PureWindowsPath(name).drive or "\\" in name:
        return False
    path = PurePosixPath(name)
    return not path.is_absolute() and ".." not in path.parts


def _validate_ooxml(data: bytes, extension: str) -> dict[str, Any]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_MEMBERS:
                return _error("archive_too_many_members", "Office archive has too many entries")
            names: set[str] = set()
            expanded = 0
            for member in infos:
                if not _safe_archive_name(member.filename):
                    return _error("unsafe_archive_path", "Office archive contains an unsafe path")
                if member.filename in names:
                    return _error("duplicate_archive_member", "Office archive contains duplicate entries")
                names.add(member.filename)
                if member.flag_bits & 0x1:
                    return _error("encrypted_document", "Password-protected Office files are not supported",
                                  status="UNSUPPORTED")
                if member.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                    return _error("archive_member_too_large", "Office archive entry exceeds the safety limit")
                expanded += max(0, member.file_size)
                if expanded > MAX_ARCHIVE_EXPANDED_BYTES:
                    return _error("archive_expands_too_large", "Office archive expands beyond the safety limit")
                if member.file_size > 1024 * 1024:
                    ratio = member.file_size / max(member.compress_size, 1)
                    if ratio > MAX_COMPRESSION_RATIO:
                        return _error("archive_compression_ratio", "Office archive has an unsafe compression ratio")
                mode = (member.external_attr >> 16) & 0o170000
                if mode == 0o120000:
                    return _error("archive_link_member", "Office archive links are not supported")
            required = {"[Content_Types].xml", _OOXML[extension][1]}
            if not required.issubset(names):
                return _error("malformed_ooxml", f"File is not a valid {extension[1:].upper()} document")
            content_types = archive.read("[Content_Types].xml")
            if _OOXML[extension][0].encode() not in content_types:
                return _error(
                    "ooxml_type_mismatch",
                    f"Office package content does not match {extension[1:].upper()}",
                )
            bad = archive.testzip()
            if bad:
                return _error("corrupt_archive", "Office archive failed its integrity check")
    except (zipfile.BadZipFile, EOFError, OSError, RuntimeError, ValueError):
        return _error("malformed_ooxml", f"File is not a valid {extension[1:].upper()} document")
    return {"status": "success"}


def validate_upload(data: bytes, filename: str, declared_content_type: str) -> dict[str, Any]:
    """Validate an uploaded document using bytes, extension, and bounded archive rules.

    The browser-provided MIME is evidence, never authority. Generic octet-stream
    is accepted only after the content and extension identify a supported format.
    Legacy compound Office files are rejected explicitly instead of entering an
    unreliable desktop-conversion path.
    """
    extension = os.path.splitext(os.path.basename(filename or ""))[1].lower()
    declared = (declared_content_type or "application/octet-stream").split(";", 1)[0].lower()
    if extension in _LEGACY_EXTS or data.startswith(_OLE_SIGNATURE):
        return _error(
            "legacy_office_unsupported",
            "Legacy .doc/.ppt/.xls files are not supported. Export the file as DOCX, PPTX, XLSX, or PDF and upload it again.",
            status="UNSUPPORTED",
        )
    if extension not in _SUPPORTED:
        return _error("unsupported_extension", "Supported files are PDF, DOCX, PPTX, XLSX, TXT, and CSV",
                      status="UNSUPPORTED")

    detected = _SUPPORTED[extension]
    if extension == ".pdf":
        if not data.startswith(b"%PDF-") or b"%%EOF" not in data[-4096:]:
            return _error("malformed_pdf", "File is not a complete PDF document")
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                return _error("encrypted_document", "Password-protected PDFs are not supported",
                              status="UNSUPPORTED")
            if not reader.pages:
                return _error("empty_pdf", "PDF has no pages")
            if len(reader.pages) > MAX_DOCUMENT_UNITS:
                return _error("document_too_long", f"PDF exceeds the {MAX_DOCUMENT_UNITS}-page limit")
        except ImportError:
            # Build-time dependency failure should be explicit, not misreported as
            # a bad founder file.
            return _error("extractor_unavailable", "PDF validation is not installed")
        except Exception:
            return _error("malformed_pdf", "File is not a readable PDF document")
    elif extension in _OOXML:
        if not data.startswith(b"PK"):
            return _error("type_mismatch", f"File bytes do not match {extension[1:].upper()}")
        checked = _validate_ooxml(data, extension)
        if checked.get("status") != "success":
            return checked
    else:
        if b"\x00" in data[:8192]:
            return _error("binary_text_file", "Text uploads cannot contain binary NUL bytes")
        try:
            data.decode("utf-8")
        except UnicodeDecodeError:
            return _error("invalid_text_encoding", "TXT and CSV uploads must be UTF-8")

    compatible = {detected, "application/octet-stream"}
    if extension in _OOXML:
        compatible.add("application/zip")
    if extension == ".csv":
        compatible.add("text/plain")
    if declared not in compatible:
        return _error(
            "declared_type_mismatch",
            f"Declared content type {declared!r} does not match detected type {detected!r}",
        )
    return {
        "status": "success", "detected_extension": extension,
        "detected_content_type": detected,
    }


def _chunk_record(content: str, ordinal: int, kind: str, locator: dict[str, Any],
                  heading: str = "") -> dict[str, Any]:
    normalized = re.sub(r"[ \t]+", " ", content).strip()
    return {
        "ordinal": ordinal,
        "content": normalized,
        "content_sha256": hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        "kind": kind,
        "heading": heading[:300],
        "locator": locator,
        "char_count": len(normalized),
    }


def _split_unit(content: str, *, ordinal: int, kind: str,
                locator: dict[str, Any], heading: str = "") -> list[dict[str, Any]]:
    content = content.strip()
    if not content:
        return []
    if len(content) <= MAX_CHUNK_CHARS:
        return [_chunk_record(content, ordinal, kind, locator, heading)]
    pieces: list[dict[str, Any]] = []
    paragraphs = re.split(r"\n\s*\n", content)
    current = ""
    part = 1
    for paragraph in paragraphs:
        candidate = f"{current}\n\n{paragraph}".strip()
        if current and len(candidate) > MAX_CHUNK_CHARS:
            pieces.append(_chunk_record(current, ordinal + len(pieces), kind,
                                        {**locator, "part": part}, heading))
            if len(pieces) >= MAX_CHUNKS:
                return pieces
            part += 1
            current = paragraph
        else:
            current = candidate
        while len(current) > MAX_CHUNK_CHARS:
            boundary = current.rfind("\n", 0, MAX_CHUNK_CHARS)
            if boundary < MAX_CHUNK_CHARS // 2:
                boundary = MAX_CHUNK_CHARS
            pieces.append(_chunk_record(current[:boundary], ordinal + len(pieces), kind,
                                        {**locator, "part": part}, heading))
            if len(pieces) >= MAX_CHUNKS:
                return pieces
            part += 1
            current = current[boundary:].lstrip()
    if current:
        pieces.append(_chunk_record(current, ordinal + len(pieces), kind,
                                    {**locator, "part": part}, heading))
    return pieces


def _extract_pdf(path: str) -> list[dict[str, Any]]:
    from pypdf import PdfReader
    reader = PdfReader(path, strict=False)
    chunks: list[dict[str, Any]] = []
    for page_no, page in enumerate(reader.pages[:MAX_DOCUMENT_UNITS], start=1):
        if len(chunks) >= MAX_CHUNKS:
            break
        # pypdf raises KeyError for a legitimate blank page with no /Contents
        # stream. That is a NO_TEXT result, not an extractor failure.
        text = (page.extract_text(extraction_mode="layout") or "") \
            if "/Contents" in page else ""
        chunks.extend(_split_unit(
            text, ordinal=len(chunks), kind="page", locator={"page": page_no},
        )[:MAX_CHUNKS - len(chunks)])
    return chunks


def _extract_docx(path: str) -> list[dict[str, Any]]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    chunks: list[dict[str, Any]] = []
    section = ""
    paragraph_no = 0
    for block in doc.iter_inner_content():
        if len(chunks) >= MAX_CHUNKS:
            break
        if isinstance(block, Paragraph):
            paragraph_no += 1
            text = block.text.strip()
            if not text:
                continue
            style = (getattr(block.style, "name", "") or "").lower()
            if style.startswith("heading"):
                section = text
            chunks.extend(_split_unit(
                text, ordinal=len(chunks), kind="heading" if style.startswith("heading") else "paragraph",
                locator={"paragraph": paragraph_no, **({"section": section} if section else {})},
                heading=section,
            )[:MAX_CHUNKS - len(chunks)])
        elif isinstance(block, Table):
            rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells]
                    for row in block.rows]
            if rows:
                width = max(len(row) for row in rows)
                rows = [row + [""] * (width - len(row)) for row in rows]
                md = ["| " + " | ".join(row) + " |" for row in rows]
                if len(md) > 1:
                    md.insert(1, "| " + " | ".join(["---"] * width) + " |")
                chunks.extend(_split_unit(
                    "\n".join(md), ordinal=len(chunks), kind="table",
                    locator={"paragraph": paragraph_no, **({"section": section} if section else {})},
                    heading=section,
                )[:MAX_CHUNKS - len(chunks)])
    # Headers and footers are not part of iter_inner_content.
    for section_no, doc_section in enumerate(doc.sections, start=1):
        if len(chunks) >= MAX_CHUNKS:
            break
        for area, paragraphs in (("header", doc_section.header.paragraphs),
                                 ("footer", doc_section.footer.paragraphs)):
            text = "\n".join(p.text.strip() for p in paragraphs if p.text.strip())
            chunks.extend(_split_unit(
                text, ordinal=len(chunks), kind=area,
                locator={"section": section_no, "area": area},
            )[:MAX_CHUNKS - len(chunks)])
    return chunks


def _pptx_chart_markdown(chart: Any) -> str:
    try:
        categories = [str(c.label) for c in chart.plots[0].categories]
        series = list(chart.series)
        rows = [["Category", *[str(s.name) for s in series]]]
        for index, category in enumerate(categories):
            rows.append([category, *[str(s.values[index]) for s in series]])
        return "\n".join(
            ["| " + " | ".join(row) + " |" for row in rows[:1]]
            + ["| " + " | ".join(["---"] * len(rows[0])) + " |"]
            + ["| " + " | ".join(row) + " |" for row in rows[1:]]
        )
    except Exception:
        return "[Unsupported chart]"


def _extract_pptx(path: str) -> list[dict[str, Any]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(path)
    chunks: list[dict[str, Any]] = []
    for slide_no, slide in enumerate(prs.slides, start=1):
        if slide_no > MAX_DOCUMENT_UNITS or len(chunks) >= MAX_CHUNKS:
            break
        title = slide.shapes.title.text.strip() if slide.shapes.title else ""
        parts: list[str] = []

        def visit(shape: Any) -> None:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for child in sorted(shape.shapes, key=lambda x: (x.top or 0, x.left or 0)):
                    visit(child)
                return
            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells]
                        for row in shape.table.rows]
                if rows:
                    width = max(len(row) for row in rows)
                    rows = [row + [""] * (width - len(row)) for row in rows]
                    parts.append("\n".join(
                        ["| " + " | ".join(rows[0]) + " |",
                         "| " + " | ".join(["---"] * width) + " |"]
                        + ["| " + " | ".join(row) + " |" for row in rows[1:]]))
            elif getattr(shape, "has_chart", False):
                parts.append(_pptx_chart_markdown(shape.chart))
            elif getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                alt = ""
                try:
                    alt = shape._element._nvXxPr.cNvPr.attrib.get("descr", "")
                except Exception:
                    pass
                parts.append(f"[Image: {alt or shape.name}]")

        for shape in sorted(slide.shapes, key=lambda x: (x.top or 0, x.left or 0)):
            visit(shape)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"### Speaker notes\n{notes}")
        chunks.extend(_split_unit(
            "\n\n".join(parts), ordinal=len(chunks), kind="slide",
            locator={"slide": slide_no}, heading=title,
        )[:MAX_CHUNKS - len(chunks)])
    return chunks


def _extract_xlsx(path: str) -> list[dict[str, Any]]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path, read_only=True, data_only=True)
    chunks: list[dict[str, Any]] = []
    for ws in wb.worksheets[:MAX_DOCUMENT_UNITS]:
        if len(chunks) >= MAX_CHUNKS:
            break
        lines: list[str] = []
        start_row = 1
        max_column = max(1, min(int(ws.max_column or 1), MAX_SHEET_COLUMNS))
        max_row = max(1, min(int(ws.max_row or 1), MAX_SHEET_ROWS))
        for row_no, row in enumerate(ws.iter_rows(
                max_row=max_row, max_col=max_column, values_only=True), start=1):
            values = ["" if value is None else str(value) for value in row]
            if not any(values):
                continue
            line = "\t".join(values)
            candidate = "\n".join([*lines, line])
            if lines and len(candidate) > MAX_CHUNK_CHARS:
                chunks.append(_chunk_record(
                    "\n".join(lines), len(chunks), "sheet",
                    {"sheet": ws.title, "cell_range": f"A{start_row}:{get_column_letter(max_column)}{row_no - 1}"},
                    ws.title,
                ))
                if len(chunks) >= MAX_CHUNKS:
                    break
                lines = [line]
                start_row = row_no
            else:
                lines.append(line)
        if lines and len(chunks) < MAX_CHUNKS:
            chunks.append(_chunk_record(
                "\n".join(lines), len(chunks), "sheet",
                {"sheet": ws.title, "cell_range": f"A{start_row}:{get_column_letter(max_column)}{max_row}"},
                ws.title,
            ))
    return chunks


def extract_chunks(path: str, detected_extension: str) -> dict[str, Any]:
    """Extract bounded format-aware chunks; errors are returned as data."""
    try:
        if detected_extension == ".pdf":
            chunks = _extract_pdf(path)
        elif detected_extension == ".docx":
            chunks = _extract_docx(path)
        elif detected_extension == ".pptx":
            chunks = _extract_pptx(path)
        elif detected_extension == ".xlsx":
            chunks = _extract_xlsx(path)
        elif detected_extension in {".txt", ".csv"}:
            with open(path, "rb") as handle:
                text = handle.read().decode("utf-8")
            chunks = _split_unit(text, ordinal=0, kind="text", locator={"section": "document"})
        else:
            return _error("unsupported_extension", "No extractor is registered for this format",
                          status="UNSUPPORTED")
    except Exception as exc:
        return _error("extraction_failed", f"Document extraction failed: {exc}"[:240])
    chunks = [chunk for chunk in chunks if chunk.get("content")][:MAX_CHUNKS]
    if not chunks:
        return _error(
            "no_extractable_text",
            "No readable text was found. If this is a scan, export it with OCR or upload a text-searchable PDF.",
            status="NO_TEXT",
        )
    return {"status": "success", "chunks": chunks, "chunk_count": len(chunks)}


async def _with_embeddings(chunks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if _embed_fn is None or not chunks:
        return chunks
    try:
        vectors: list[list[float]] = []
        texts = [c["content"][:8000] for c in chunks]
        for start in range(0, len(texts), 32):
            vectors.extend(await asyncio.to_thread(_embed_fn, texts[start:start + 32]))
        if len(vectors) == len(chunks):
            return [{**chunk, "embedding": vector} for chunk, vector in zip(chunks, vectors)]
    except Exception:
        pass
    return chunks


def _citation_for(chunk: dict[str, Any], artifact_id: str, quote: str | None = None) -> dict[str, Any]:
    selected = (quote or chunk.get("content") or "")[:1200]
    return {
        "artifact_id": artifact_id,
        "chunk_id": chunk.get("id", ""),
        "locator": chunk.get("locator") or {},
        "quote": selected,
        "quote_sha256": hashlib.sha256(selected.encode("utf-8")).hexdigest(),
    }


def _tokens(value: str) -> set[str]:
    return {token for token in re.findall(r"[a-z0-9]{2,}", value.lower())
            if token not in {"the", "and", "for", "with", "from", "that", "this"}}


def _cosine(a: list[float], b: list[float]) -> float:
    if not a or not b or len(a) != len(b):
        return 0.0
    dot = sum(x * y for x, y in zip(a, b))
    norm = math.sqrt(sum(x * x for x in a)) * math.sqrt(sum(y * y for y in b))
    return dot / norm if norm else 0.0


async def search_attachments(*, founder_id: str, session_id: str,
                             attachment_refs: list[str], query: str,
                             top_k: int = 5) -> dict[str, Any]:
    """Search explicitly referenced, owner/session-authorized artifacts."""
    query = query.strip()[:500]
    if not query:
        return _error("empty_query", "Attachment search requires a query")
    refs = list(dict.fromkeys(attachment_refs))[:8]
    if not refs:
        return _error("no_attachments", "No attachment references were provided")

    query_vector: list[float] = []
    if _embed_fn is not None:
        try:
            vectors = await asyncio.to_thread(_embed_fn, [query])
            query_vector = vectors[0] if vectors else []
        except Exception:
            query_vector = []
    query_terms = _tokens(query)
    ranked: list[tuple[float, dict[str, Any], dict[str, Any]]] = []
    for ref in refs:
        artifact = await firestore.get_artifact(ref)
        if (not artifact or artifact.get("founder_id") != founder_id
                or artifact.get("session_id") != session_id):
            return _error("attachment_not_found", "Attachment was not found")
        if artifact.get("status") not in {"READY", "NEEDS_FOUNDER", "CONFIRMED"}:
            return _error("attachment_not_ready", f"Attachment {ref} is not ready for search")
        chunks = await firestore.list_artifact_chunks(ref, limit=MAX_CHUNKS)
        for chunk in chunks:
            terms = _tokens(chunk.get("content", ""))
            lexical = len(query_terms & terms) / max(len(query_terms), 1)
            semantic = _cosine(query_vector, chunk.get("embedding") or [])
            score = max(lexical, semantic) + (0.15 * lexical if semantic else 0.0)
            if score > 0:
                ranked.append((score, artifact, chunk))
    ranked.sort(key=lambda item: (-item[0], item[2].get("ordinal", 0)))
    results = []
    for score, artifact, chunk in ranked[:max(1, min(top_k, 8))]:
        citation = _citation_for(chunk, artifact["id"])
        locator = citation.get("locator")
        if (not citation.get("artifact_id") or not citation.get("chunk_id")
                or not isinstance(locator, dict) or not locator
                or not citation.get("quote")
                or citation.get("quote_sha256") != hashlib.sha256(
                    citation["quote"].encode()).hexdigest()):
            return _error(
                "incomplete_evidence_citation",
                "Attachment evidence is missing an exact source locator; no result was returned.")
        location_key = next((key for key in (
            "page", "slide", "paragraph", "section", "sheet", "cell_range")
                             if locator.get(key) not in (None, "")), "")
        if not location_key:
            return _error(
                "incomplete_evidence_citation",
                "Attachment evidence is missing an exact source locator; no result was returned.")
        source_available = True
        if artifact.get("source_grant_id"):
            grant = await firestore.get_source_grant(
                founder_id, artifact["source_grant_id"])
            source_available = bool(grant and grant.get("status") == "ACTIVE")
        citation["source_url"] = (
            f"/api/ingest/{artifact['id']}/source?session_id={session_id}"
            f"#{location_key}={locator[location_key]}")
        results.append({
            "source_type": artifact.get("source_type") or "upload",
            "source_id": artifact["id"],
            "source_title": artifact.get("source_ref", "document"),
            "score": round(score, 4),
            "excerpt": chunk.get("content", "")[:2000],
            "citation": citation,
            "authority": "unconfirmed_evidence",
            "source_available": source_available,
        })
    return {"status": "success", "query": query, "results": results}


async def _transient_failure(ingestion_id: str, owner: str, code: str,
                             message: str, retry_transient: bool) -> dict[str, Any]:
    if retry_transient:
        released = await firestore.retry_ingestion(
            ingestion_id, owner, error_code=code, message=message)
        if released.get("status") == "success" and released.get("retryable"):
            return {**_error(code, message), "retryable": True,
                    "ingestion_status": "QUEUED"}
        if released.get("status") != "success":
            return {**released, "retryable": True,
                    "error_code": released.get("error_code", "lease_changed")}
    await firestore.finish_ingestion(
        ingestion_id, owner, "FAILED", error_code=code, message=message)
    return _error(code, message)


async def process_ingestion(ingestion_id: str, *, lease_owner: str = "",
                            retry_transient: bool = False) -> dict[str, Any]:
    """Run one idempotent ingestion attempt from persisted artifact state."""
    from services import profile_service

    claim = await firestore.claim_ingestion(ingestion_id, lease_owner=lease_owner)
    if claim.get("duplicate"):
        return {"status": "success", "duplicate": True,
                "ingestion_status": claim.get("ingestion_status")}
    if claim.get("in_progress"):
        return {"status": "success", "in_progress": True}
    if claim.get("status") != "success":
        return claim
    owner = claim["lease_owner"]
    ingestion = await firestore.get_ingestion(ingestion_id)
    artifact = await firestore.get_artifact(ingestion_id)
    if not ingestion or not artifact:
        return _error("ingestion_not_found", "Ingestion was not found")
    try:
        source_bytes = await asyncio.to_thread(storage.read_bytes, artifact["storage_name"])
        if hashlib.sha256(source_bytes).hexdigest() != artifact.get("sha256"):
            failed = _error("artifact_hash_mismatch", "Stored artifact failed its integrity check")
            await firestore.finish_ingestion(
                ingestion_id, owner, "FAILED", error_code=failed["error_code"],
                message=failed["message"])
            return failed
        checked = await asyncio.to_thread(
            validate_upload, source_bytes, artifact.get("source_ref", ""),
            artifact.get("declared_content_type", "application/octet-stream"))
        if checked.get("status") != "success":
            await firestore.finish_ingestion(
                ingestion_id, owner, checked["ingestion_status"],
                error_code=checked.get("error_code"), message=checked.get("message"))
            return checked
        if not await firestore.set_ingestion_stage(ingestion_id, owner, "EXTRACTING"):
            return {**_error("lease_changed", "Document worker lease changed"),
                    "retryable": True}
        path = storage.download_if_missing(artifact["storage_name"])
        extracted = await asyncio.to_thread(
            extract_chunks, path, artifact["detected_extension"])
        if extracted.get("status") != "success":
            await firestore.finish_ingestion(
                ingestion_id, owner, extracted["ingestion_status"],
                error_code=extracted.get("error_code", "extraction_failed"),
                message=extracted.get("message", "Document extraction failed"),
            )
            return extracted

        if not await firestore.set_ingestion_stage(ingestion_id, owner, "INDEXING"):
            return {**_error("lease_changed", "Document worker lease changed"),
                    "retryable": True}
        chunks = await _with_embeddings(extracted["chunks"])
        await firestore.replace_artifact_chunks(ingestion_id, chunks)

        proposed_count = auto_applied = needs_founder = 0
        final_status = "READY"
        if artifact.get("scope") == "profile":
            staged = ingestion.get("proposed_updates") or []
            proposed = ({"status": "success", "proposals": staged} if staged else
                        await profile_service.extract_profile_proposals(
                            artifact["founder_id"], ingestion_id,
                            artifact["storage_name"], chunks))
            if proposed.get("status") != "success":
                code = proposed.get("error_code", "proposal_extraction_failed")
                message = proposed.get("message", "Profile extraction failed")
                if code == "proposal_extraction_failed":
                    return await _transient_failure(
                        ingestion_id, owner, code, message, retry_transient)
                await firestore.finish_ingestion(
                    ingestion_id, owner, "FAILED", error_code=code, message=message)
                return proposed
            proposals = proposed["proposals"]
            proposed_count = len(proposals)
            staged_ok = await firestore.update_ingestion_leased(
                ingestion_id, owner, proposed_updates=proposals)
            if not staged_ok:
                return {**_error("lease_changed", "Document worker lease changed"),
                        "retryable": True}
            reviewed = await profile_service.auto_apply_profile_updates(
                artifact["founder_id"], ingestion_id, commit_status=False,
                lease_owner=owner)
            if reviewed.get("status") != "success":
                if reviewed.get("error_code") == "lease_changed":
                    return {**reviewed, "retryable": True}
                await firestore.finish_ingestion(
                    ingestion_id, owner, "FAILED", error_code="profile_apply_failed",
                    message=reviewed.get("message", "Profile update failed"),
                )
                return reviewed
            auto_applied = int(reviewed.get("auto_applied") or 0)
            needs_founder = len(reviewed.get("needs_founder") or [])
            final_status = "NEEDS_FOUNDER" if needs_founder else "CONFIRMED"

        committed = await firestore.finish_ingestion(
            ingestion_id, owner, final_status, chunk_count=len(chunks),
            proposed_count=proposed_count, auto_applied=auto_applied,
            needs_founder_count=needs_founder,
        )
        if not committed:
            return {**_error("lease_changed", "Document result was not committed"),
                    "retryable": True}
        return {
            "status": "success", "attachment_ref": ingestion_id,
            "ingestion_status": final_status, "chunk_count": len(chunks),
            "proposed_count": proposed_count, "auto_applied": auto_applied,
            "needs_founder": needs_founder,
        }
    except Exception as exc:
        return await _transient_failure(
            ingestion_id, owner, "worker_failure",
            f"Document processing failed: {exc}"[:240], retry_transient)
